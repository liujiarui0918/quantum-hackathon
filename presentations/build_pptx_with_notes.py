from __future__ import annotations

import re
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PRESENTATIONS = ROOT / "presentations"
PDF_PATH = PRESENTATIONS / "group_meeting_qbb_miqp_2026_05_21.pdf"
SPEECH_PATH = PRESENTATIONS / "group_meeting_qbb_miqp_2026_05_21_speech.md"
FRAMES_DIR = PRESENTATIONS / "pptx_frames"
OUT_PATH = PRESENTATIONS / "group_meeting_qbb_miqp_2026_05_21_with_notes.pptx"

SLIDE_W = Inches(13.333333)
SLIDE_H = Inches(7.5)


def run(command: list[str]) -> str:
    completed = subprocess.run(
        command,
        cwd=ROOT,
        check=True,
        encoding="utf-8",
        errors="replace",
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
    )
    return completed.stdout


def pdf_page_count(pdf_path: Path) -> int:
    output = run(["pdfinfo", str(pdf_path)])
    match = re.search(r"^Pages:\s+(\d+)", output, re.MULTILINE)
    if not match:
        raise RuntimeError("Could not read PDF page count.")
    return int(match.group(1))


def clean_frame_dir() -> None:
    FRAMES_DIR.mkdir(exist_ok=True)
    for image in FRAMES_DIR.glob("slide-*.png"):
        image.unlink()


def render_pdf_pages() -> list[Path]:
    clean_frame_dir()
    run(
        [
            "pdftoppm",
            "-png",
            "-r",
            "300",
            str(PDF_PATH),
            str(FRAMES_DIR / "slide"),
        ]
    )
    return sorted(FRAMES_DIR.glob("slide-*.png"))


def clean_note_text(text: str) -> str:
    text = re.sub(r"^#+\s*", "", text, flags=re.MULTILINE)
    text = text.replace("  \n", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_speaker_notes() -> dict[int, str]:
    text = SPEECH_PATH.read_text(encoding="utf-8")
    notes: dict[int, str] = {}
    matches = list(re.finditer(r"^## 第\s+(\d+)\s+页：(.+)$", text, re.MULTILINE))
    for index, match in enumerate(matches):
        page = int(match.group(1))
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        block = text[start:end]
        notes[page] = clean_note_text(block)

    appendix_match = re.search(r"^## 附录页讲法\s*$", text, re.MULTILINE)
    if appendix_match:
        appendix = text[appendix_match.start() :]
        appendix_intro = re.search(
            r"^## 附录页讲法\s*(.*?)(?=^###\s|\Z)",
            appendix,
            re.MULTILINE | re.DOTALL,
        )
        if appendix_intro:
            notes[57] = clean_note_text("附录页讲法\n" + appendix_intro.group(1))

        appendix_sections = list(
            re.finditer(r"^###\s+(.+)$", appendix, re.MULTILINE)
        )
        for offset, section in enumerate(appendix_sections, start=58):
            start = section.start()
            end = (
                appendix_sections[offset - 57].start()
                if offset - 57 < len(appendix_sections)
                else len(appendix)
            )
            notes[offset] = clean_note_text(appendix[start:end])

    notes.setdefault(
        62,
        "参考材料\n\n这一页作为备查材料保留。报告时一般不展开讲，只在需要说明材料来源、学习讲义或代码位置时引用。",
    )
    return notes


def add_notes(slide, note_text: str) -> None:
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    paragraphs = [part.strip() for part in note_text.split("\n\n") if part.strip()]
    if not paragraphs:
        paragraphs = [""]
    notes_frame.text = paragraphs[0]
    for paragraph in paragraphs[1:]:
        p = notes_frame.add_paragraph()
        p.text = paragraph
        p.font.size = Pt(12)


def build_pptx(frame_paths: list[Path], notes: dict[int, str]) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]

    for page, frame_path in enumerate(frame_paths, start=1):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(str(frame_path), 0, 0, width=SLIDE_W, height=SLIDE_H)
        add_notes(slide, notes.get(page, f"第 {page} 页\n\n"))

    prs.save(OUT_PATH)


def main() -> None:
    expected_pages = pdf_page_count(PDF_PATH)
    frame_paths = render_pdf_pages()
    if len(frame_paths) != expected_pages:
        raise RuntimeError(
            f"Rendered {len(frame_paths)} images, expected {expected_pages}."
        )
    notes = parse_speaker_notes()
    build_pptx(frame_paths, notes)
    print(f"Wrote {OUT_PATH}")
    print(f"Slides: {len(frame_paths)}")
    print(f"Notes pages: {len([p for p in range(1, len(frame_paths) + 1) if notes.get(p)])}")


if __name__ == "__main__":
    main()
